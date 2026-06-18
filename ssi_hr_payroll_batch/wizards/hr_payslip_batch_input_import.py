# Copyright 2022 OpenSynergy Indonesia
# Copyright 2022 PT. Simetri Sinergi Indonesia
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl-3.0-standalone.html).

import base64
import io

from odoo import _, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class HrPayslipBatchInputImport(models.TransientModel):
    _name = "hr.payslip_batch_input_import"
    _description = "Import Payslip Batch Input"

    batch_id = fields.Many2one(
        string="Payslip Batch",
        comodel_name="hr.payslip_batch",
        required=True,
        readonly=True,
        default=lambda self: self.env.context.get("active_id"),
    )
    data = fields.Binary(
        string="File",
        required=True,
        help="Spreadsheet (.xlsx) previously exported from this payslip batch, "
        "with the manual input amounts filled in.",
    )
    filename = fields.Char(
        string="Filename",
    )

    def _read_rows(self):
        """Decode the uploaded xlsx file and return its rows as a list of
        tuples (header row first)."""
        self.ensure_one()
        if load_workbook is None:
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: Python library 'openpyxl' is not available
Solution: Install the 'openpyxl' library on the Odoo server
"""
                % (self.id)
            )
            raise UserError(error_message)
        try:
            content = base64.b64decode(self.data)
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as error:
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: The uploaded file is not a valid xlsx spreadsheet (%s)
Solution: Upload the xlsx file exported from this payslip batch
"""
                % (self.id, error)
            )
            raise UserError(error_message)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not rows:
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: The uploaded spreadsheet is empty
Solution: Upload the xlsx file exported from this payslip batch
"""
                % (self.id)
            )
            raise UserError(error_message)
        return rows

    def _get_code_columns(self, header):
        """Map each dynamic input column index to its input code. The first
        two columns ('Payslip ID' and 'Employee') are skipped."""
        code_columns = {}
        for index, value in enumerate(header):
            if index <= 1 or not value:
                continue
            code_columns[index] = str(value).strip()
        return code_columns

    def _check_input_codes(self, codes):
        self.ensure_one()
        obj_input_type = self.env["hr.payslip_input_type"]
        unknown_codes = []
        for code in codes:
            if not obj_input_type.search([("code", "=", code)], limit=1):
                unknown_codes.append(code)
        if unknown_codes:
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: Unknown input code(s) in spreadsheet header: %s
Solution: Use the xlsx file exported from this payslip batch without changing
the header row
"""
                % (self.id, ", ".join(unknown_codes))
            )
            raise UserError(error_message)

    def _get_payslip(self, payslip_id):
        self.ensure_one()
        payslip = self.batch_id.payslip_ids.filtered(
            lambda record: record.id == payslip_id
        )
        if not payslip:
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: Payslip ID %s does not belong to this payslip batch
Solution: Use the xlsx file exported from this payslip batch
"""
                % (self.id, payslip_id)
            )
            raise UserError(error_message)
        if payslip.state != "draft":
            error_message = _(
                """
Context: Import payslip batch input
Database ID: %s
Problem: Payslip %s is not in Draft state and cannot be edited
Solution: Only import inputs while every payslip is still in Draft state
"""
                % (self.id, payslip.display_name)
            )
            raise UserError(error_message)
        return payslip

    def action_import(self):
        self.ensure_one()
        rows = self._read_rows()
        header = rows[0]
        code_columns = self._get_code_columns(header)
        self._check_input_codes(code_columns.values())

        for row in rows[1:]:
            if not row or row[0] in (None, ""):
                continue
            try:
                payslip_id = int(row[0])
            except (TypeError, ValueError):
                continue
            payslip = self._get_payslip(payslip_id)
            for index, code in code_columns.items():
                if index >= len(row):
                    continue
                value = row[index]
                if value in (None, ""):
                    continue
                input_line = payslip.input_line_ids.filtered(
                    lambda record: record.input_type_id.code == code
                )
                if input_line:
                    input_line.write({"amount": float(value)})
        return {"type": "ir.actions.act_window_close"}
