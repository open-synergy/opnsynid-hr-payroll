# Copyright 2025 OpenSynergy Indonesia
# Copyright 2025 PT. Simetri Sinergi Indonesia
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).

import base64
import io

import xlsxwriter
from odoo_yaml_test import YamlTransactionCase
from openpyxl import load_workbook

from odoo.tests import tagged


@tagged("post_install", "-at_install")
class TestHrPayslipBatch(YamlTransactionCase):
    """Tests ``hr.payslip_batch`` workflow, payslip counts and inputs."""

    def test_hr_payslip_batch(self):
        """Runs the YAML batch workflow, onchange, count and import cases."""
        self.run_yaml_scenario("test_data_hr_payslip_batch.yaml")

    def test_prepare_payslip_data_includes_analytic_account(self):
        """``_prepare_payslip_data`` must include ``analytic_account_id``.

        Pure Python -- trigger P1 (L-01: ``action: call`` discards a
        method's return value; L-02: an ``assert`` target must be a
        stored record field, not an ad-hoc dict).
        """
        analytic = self.env["account.analytic.account"].create(
            {"name": "Test Batch Propagation Analytic"}
        )
        journal = self.env["account.journal"].create(
            {
                "name": "Test Batch Propagation Journal",
                "code": "TBPRJRN",
                "type": "general",
            }
        )
        rule_cat = self.env["hr.salary_rule_category"].create(
            {"name": "Test Batch Prop Cat", "code": "TBPROPCAT"}
        )
        rule = self.env["hr.salary_rule"].create(
            {
                "name": "Test Batch Prop Rule",
                "code": "TBPROPRULE",
                "category_id": rule_cat.id,
                "condition_python": "result = True",
                "amount_python": "result = 100.0",
                "sequence": 10,
            }
        )
        structure = self.env["hr.salary_structure"].create(
            {
                "name": "Test Batch Prop Structure",
                "code": "TBPROPSTR",
                "rule_ids": [(4, rule.id)],
            }
        )
        payslip_type = self.env["hr.payslip_type"].create(
            {
                "name": "Test Batch Prop Type",
                "code": "TBPROPTYPE",
                "journal_id": journal.id,
                "analytic_account_id": analytic.id,
            }
        )
        employee = self.env["hr.employee"].create(
            {
                "name": "Test Batch Prop Employee",
                "salary_structure_id": structure.id,
            }
        )
        batch = self.env["hr.payslip_batch"].create(
            {
                "type_id": payslip_type.id,
                "analytic_account_id": analytic.id,
                "date_start": "2024-11-01",
                "date_end": "2024-11-30",
                "date": "2024-11-30",
            }
        )
        payslip_data = batch._prepare_payslip_data(employee)
        self.assertEqual(
            payslip_data.get("analytic_account_id"),
            analytic.id,
        )

    def _create_batch_with_input(self):
        """Build a payslip batch whose generated payslips carry two manual
        input types (BONUS and OVERTIME). Returns (batch, input_types)."""
        bonus_type = self.env["hr.payslip_input_type"].create(
            {"name": "Test Bonus Input", "code": "BONUS"}
        )
        overtime_type = self.env["hr.payslip_input_type"].create(
            {"name": "Test Overtime Input", "code": "OVERTIME"}
        )
        journal = self.env["account.journal"].create(
            {
                "name": "Test Input Journal",
                "code": "TINPJRN",
                "type": "general",
            }
        )
        rule_cat = self.env["hr.salary_rule_category"].create(
            {"name": "Test Input Cat", "code": "TINPCAT"}
        )
        rule = self.env["hr.salary_rule"].create(
            {
                "name": "Test Input Rule",
                "code": "TINPRULE",
                "category_id": rule_cat.id,
                "condition_python": "result = True",
                "amount_python": "result = 100.0",
                "sequence": 10,
                "input_type_ids": [(6, 0, [bonus_type.id, overtime_type.id])],
            }
        )
        structure = self.env["hr.salary_structure"].create(
            {
                "name": "Test Input Structure",
                "code": "TINPSTR",
                "rule_ids": [(4, rule.id)],
            }
        )
        payslip_type = self.env["hr.payslip_type"].create(
            {
                "name": "Test Input Type",
                "code": "TINPTYPE",
                "journal_id": journal.id,
            }
        )
        emp_vals = {"name": "Test Input Employee"}
        emp_model = self.env["hr.employee"]
        if "method" in emp_model._fields:
            emp_vals["method"] = "manual"
        struct_field = (
            "manual_salary_structure_id"
            if "manual_salary_structure_id" in emp_model._fields
            else "salary_structure_id"
        )
        emp_vals[struct_field] = structure.id
        employee = emp_model.create(emp_vals)
        batch = self.env["hr.payslip_batch"].create(
            {
                "type_id": payslip_type.id,
                "date_start": "2024-11-01",
                "date_end": "2024-11-30",
                "date": "2024-11-30",
                "employee_ids": [(6, 0, [employee.id])],
            }
        )
        batch._generate_payslip()
        return batch, (bonus_type + overtime_type)

    def _read_attachment_rows(self, url):
        """Decode the xlsx attachment referenced by an act_url and return
        its rows as a list of tuples."""
        attachment_id = int(url.split("/web/content/")[1].split("?")[0])
        attachment = self.env["ir.attachment"].browse(attachment_id)
        content = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        return rows

    def test_action_export_input_contains_dynamic_columns(self):
        """Export must produce one column per input code present in the
        batch, plus the fixed Payslip ID and Employee columns.

        Pure Python -- trigger P4 (L-08: the content of a binary/
        attachment field cannot be asserted from YAML; decoding the
        exported xlsx requires reading its binary payload).
        """
        batch, _input_types = self._create_batch_with_input()
        payslip = batch.payslip_ids
        self.assertEqual(len(payslip), 1)

        action = batch.action_export_input()
        self.assertEqual(action["type"], "ir.actions.act_url")

        rows = self._read_attachment_rows(action["url"])
        header = list(rows[0])
        self.assertEqual(header[0], "Payslip ID")
        self.assertEqual(header[1], "Employee")
        self.assertIn("BONUS", header)
        self.assertIn("OVERTIME", header)

        # One data row per payslip, with matching id and employee name.
        self.assertEqual(len(rows) - 1, len(batch.payslip_ids))
        data_row = rows[1]
        self.assertEqual(int(data_row[0]), payslip.id)
        self.assertEqual(data_row[1], payslip.employee_id.name)

    def test_input_import_updates_amounts(self):
        """Importing a spreadsheet must write the input amounts back onto
        the matching payslip input lines.

        Pure Python -- trigger P10 (L-09/L-10/L-11: the ``EVAL:``
        sandbox has no ``import`` and no ``xlsxwriter``/``base64``, so
        the uploaded xlsx fixture -- keyed by this run's dynamically
        assigned payslip ID -- cannot be built from YAML).
        """
        batch, input_types = self._create_batch_with_input()
        payslip = batch.payslip_ids

        # Build an xlsx mirroring the export layout with amounts filled in.
        codes = [it.code for it in batch._get_batch_input_type_ids()]
        amounts = {"BONUS": 150000.0, "OVERTIME": 75000.0}
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        sheet = workbook.add_worksheet("Payslip Inputs")
        for column, header in enumerate(["Payslip ID", "Employee"] + codes):
            sheet.write(0, column, header)
        sheet.write_number(1, 0, payslip.id)
        sheet.write_string(1, 1, payslip.employee_id.name or "")
        for column, code in enumerate(codes, start=2):
            sheet.write_number(1, column, amounts[code])
        workbook.close()
        output.seek(0)

        wizard = self.env["import_payslip_batch_input"].create(
            {
                "batch_id": batch.id,
                "data": base64.b64encode(output.read()),
                "filename": "input.xlsx",
            }
        )
        wizard.action_import()

        for input_type in input_types:
            line = payslip.input_line_ids.filtered(
                lambda record, t=input_type: record.input_type_id == t
            )
            self.assertEqual(line.amount, amounts[input_type.code])

    def test_action_open_payslip(self):
        """``action_open_payslip`` must return a window action scoped to
        this batch's payslips.

        Pure Python -- trigger P1 (L-01: ``action: call`` discards the
        returned ``ir.actions.act_window`` dict; L-02: an ``assert``
        target must be a stored record field, not the returned dict).
        """
        batch, _input_types = self._create_batch_with_input()
        action = batch.action_open_payslip()
        self.assertEqual(action["res_model"], "hr.payslip")
        self.assertEqual(action["domain"], [("batch_id", "=", batch.id)])
        self.assertEqual(
            self.env["hr.payslip"].search(action["domain"]),
            batch.payslip_ids,
        )
